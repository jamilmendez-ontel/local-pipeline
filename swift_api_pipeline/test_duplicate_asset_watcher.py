"""Unit tests for duplicate_asset_watcher's pure helpers (no network, no DB).
Run: cd swift_api_pipeline && python -m pytest test_duplicate_asset_watcher.py -v
"""

from datetime import datetime, timezone

from duplicate_asset_watcher import (
    push_created_at, group_rows, classify_group, select_groups, is_full_list_day,
    build_email_body, build_email_html, build_workbook, build_message, attach_blank_timers,
    sort_groups, section_of, summarize, NEW_HOURS,
)

KEEP = "-P3sPO9nXQmaGuy_lLm8"    # created 2026-10-14 03:16:45 UTC, 5 approved tasks, 2 timers
DUP = "-P3sPXeBfpUeNzcPFNhW"     # created 39 s later, all 76 tasks pending, no timers


def _row(did, touched=0, timers=0, name="WINCHESTER - 5G L-SUB6 - CARRIER ADD",
         project="TECH-OPS: TS20", path="MRC Towers/VZW/Mountain Plains-KS/Embedded/17433804/Oct 2026"):
    return {"project_name": project, "project_did": "-Ozh7xx_YHk8yXEmVmFR", "asset_id": path,
            "asset_name": name, "asset_did": did, "task_count": 76, "touched": touched,
            "timers": timers}


def test_push_created_at_decodes_the_swift_id_timestamp():
    assert push_created_at(KEEP) == datetime(2026, 10, 14, 3, 16, 45, 363000, tzinfo=timezone.utc)
    assert push_created_at(DUP) == datetime(2026, 10, 14, 3, 17, 24, 236000, tzinfo=timezone.utc)


def test_push_created_at_returns_none_for_a_non_push_id():
    assert push_created_at("") is None
    assert push_created_at("not a push id!") is None


def test_group_rows_groups_by_project_path_and_name():
    rows = [_row(KEEP, touched=13, timers=2), _row(DUP),
            _row("-Pa", name="OTHER"), _row("-Pb", name="OTHER")]
    groups = group_rows(rows)
    assert len(groups) == 2
    assert {a["asset_did"] for a in groups[0]["assets"]} == {KEEP, DUP}


def test_one_active_copy_is_kept_and_the_empty_ones_are_flagged():
    g = classify_group([_row(KEEP, touched=13, timers=2), _row(DUP)])
    assert g["verdict"] == "delete_empty"
    assert g["keep"] == KEEP and g["remove"] == [DUP]


def test_all_empty_copies_keep_the_oldest():
    g = classify_group([_row(DUP), _row(KEEP)])
    assert g["verdict"] == "all_empty"
    assert g["keep"] == KEEP and g["remove"] == [DUP]


def test_two_active_copies_need_a_merge_and_nothing_is_flagged_for_removal():
    g = classify_group([_row(KEEP, touched=5), _row(DUP, timers=1)])
    assert g["verdict"] == "merge_needed"
    assert g["keep"] is None and g["remove"] == []


def test_daily_run_reports_only_groups_with_a_copy_created_recently():
    old = classify_group([_row(KEEP, touched=1), _row(DUP)])
    now = datetime(2026, 10, 14, 12, 0, tzinfo=timezone.utc)
    assert select_groups([old], now, full=False) == [old]
    later = datetime(2026, 10, 14, 3, 17, 25, tzinfo=timezone.utc).replace(day=16)
    assert (later - push_created_at(DUP)).total_seconds() > NEW_HOURS * 3600
    assert select_groups([old], later, full=False) == []
    assert select_groups([old], later, full=True) == [old]


def test_full_list_goes_out_on_monday_eastern_time():
    # 2026-10-19 is a Monday; 03:00 UTC is still Sunday evening ET
    assert is_full_list_day(datetime(2026, 10, 19, 13, 0, tzinfo=timezone.utc))
    assert not is_full_list_day(datetime(2026, 10, 19, 3, 0, tzinfo=timezone.utc))
    assert not is_full_list_day(datetime(2026, 10, 20, 13, 0, tzinfo=timezone.utc))


def test_email_body_names_the_copy_to_keep_and_the_one_to_delete():
    g = classify_group([_row(KEEP, touched=13, timers=2), _row(DUP)])
    body = build_email_body([g], datetime(2026, 10, 19, 13, 0, tzinfo=timezone.utc), full=True, total=1)
    assert "TECH-OPS: TS20" in body and "WINCHESTER - 5G L-SUB6 - CARRIER ADD" in body
    keep_line = next(l for l in body.splitlines() if KEEP in l)
    dup_line = next(l for l in body.splitlines() if DUP in l)
    assert "KEEP" in keep_line and "13 of 76 tasks touched" in keep_line and "2 timers" in keep_line
    assert "DELETE" in dup_line
    assert "—" not in body   # house style: no em dash in copy


def test_email_body_for_a_merge_case_says_so():
    g = classify_group([_row(KEEP, touched=5), _row(DUP, timers=1)])
    body = build_email_body([g], datetime(2026, 10, 19, 13, 0, tzinfo=timezone.utc), full=False, total=1)
    assert "MERGE" in body


def test_a_copy_with_unknown_task_state_is_never_flagged_for_deletion():
    g = classify_group([_row(KEEP, touched=13, timers=2), _row(DUP, touched=None)])
    assert g["verdict"] == "merge_needed" and g["remove"] == []
    body = build_email_body([g], datetime(2026, 10, 19, 13, 0, tzinfo=timezone.utc), full=True, total=1)
    assert "task state unknown" in body and "None" not in body


def test_timer_counts_are_aggregated_once_not_per_asset():
    from duplicate_asset_watcher import DUPLICATES_SQL
    assert "tc AS (" in DUPLICATES_SQL and "LEFT JOIN tc ON tc.asset_did = s.asset_did" in DUPLICATES_SQL
    assert "(SELECT count(*)" not in DUPLICATES_SQL


NOW = datetime(2026, 10, 19, 13, 0, tzinfo=timezone.utc)


def _blank(name="WINCHESTER - 5G L-SUB6 - CARRIER ADD", minutes=15.0, day=15):
    return {"project_did": "-Ozh7xx_YHk8yXEmVmFR",
            "asset_id": "MRC Towers/VZW/Mountain Plains-KS/Embedded/17433804/Oct 2026",
            "asset_name": name, "user_email": "member@ontel.co",
            "start_time": datetime(2026, 10, day, 16, 0, tzinfo=timezone.utc),
            "end_time": datetime(2026, 10, day, 16, 15, tzinfo=timezone.utc),
            "duration_min": minutes, "task": "2. Site Set-Up Complete"}


def _groups():
    hit = classify_group([_row(KEEP, touched=13, timers=2), _row(DUP)])
    merge = classify_group([_row("-Pm1", touched=5, name="MERGE ME"), _row("-Pm2", timers=1, name="MERGE ME")])
    clutter = classify_group([_row("-Pc1", name="CLUTTER"), _row("-Pc2", name="CLUTTER")])
    return attach_blank_timers([clutter, merge, hit], [_blank(minutes=15.0), _blank(minutes=45.0, day=16),
                                                        _blank(name="NOT A DUPLICATE")])


def test_blank_timers_attach_to_their_group_with_count_hours_and_last_start():
    groups = {g["asset_name"]: g for g in _groups()}
    hit = groups["WINCHESTER - 5G L-SUB6 - CARRIER ADD"]
    assert hit["blank_count"] == 2 and hit["blank_hours"] == 1.0
    assert hit["last_blank"] == datetime(2026, 10, 16, 16, 0, tzinfo=timezone.utc)
    assert groups["CLUTTER"]["blank_count"] == 0 and groups["CLUTTER"]["last_blank"] is None


def test_groups_sort_hides_timers_first_then_merge_then_clutter():
    ordered = sort_groups(_groups())
    assert [section_of(g) for g in ordered] == [1, 2, 3]
    assert ordered[0]["asset_name"].startswith("WINCHESTER")


def test_summary_counts_affected_groups_timers_hours_and_merges():
    s = summarize(_groups())
    assert s == {"groups": 3, "affected": 1, "timers": 2, "hours": 1.0, "merge": 1, "copies_to_delete": 2}


def test_html_email_has_kpis_sections_in_order_and_escapes_text():
    groups = sort_groups(_groups())
    groups[2]["asset_name"] = "A <b>&"
    body = build_email_html(groups, NOW, full=True, total=3)
    assert "Duplicate Swift assets" in body and "Blank hours" in body
    i1, i2, i3 = (body.index(t) for t in ("Fix first", "Needs a merge", "Clutter only"))
    assert i1 < i2 < i3
    assert "A &lt;b&gt;&amp;" in body and "A <b>&" not in body
    assert KEEP in body and DUP in body and "DELETE" in body and "HAS WORK" in body
    assert "—" not in body


def test_html_email_skips_empty_sections():
    only_clutter = [g for g in _groups() if g["asset_name"] == "CLUTTER"]
    body = build_email_html(only_clutter, NOW, full=False, total=3)
    assert "Clutter only" in body and "Fix first" not in body and "Needs a merge" not in body


def test_workbook_has_four_sheets_with_one_row_per_group_copy_and_blank_timer():
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(build_workbook(sort_groups(_groups()), NOW, True, 3)))
    assert wb.sheetnames == ["Summary", "Groups", "Copies", "Blank timers"]
    assert wb["Groups"].max_row == 1 + 3
    assert wb["Copies"].max_row == 1 + 6
    assert wb["Blank timers"].max_row == 1 + 2
    g = wb["Groups"]
    assert g["A2"].value.startswith("1") and g["F2"].value == 2 and g["I2"].value == KEEP and g["J2"].value == DUP
    assert g["A2"].font.name == "Arial"


def test_message_is_html_with_text_fallback_and_the_workbook_attached():
    msg = build_message("subj", "plain", "<b>x</b>", b"PK", "dup.xlsx")
    types = [p.get_content_type() for p in msg.walk()]
    assert types[:4] == ["multipart/mixed", "multipart/alternative", "text/plain", "text/html"]
    att = list(msg.walk())[-1]
    assert att.get_filename() == "dup.xlsx" and att.get_content_type().endswith("spreadsheetml.sheet")


def test_totals_round_once_from_raw_minutes_not_from_rounded_groups():
    # three groups of 2.04 h each: per-group rounding gives 2.0 x 3 = 6.0, the truth is 6.1
    from duplicate_asset_watcher import by_project
    groups = []
    for n in ("A", "B", "C"):
        g = classify_group([_row("-Pa" + n, touched=1, name=n), _row("-Pb" + n, name=n)])
        groups.append(g)
    blanks = [_blank(name=n, minutes=122.4) for n in ("A", "B", "C")]
    attach_blank_timers(groups, blanks)
    assert [g["blank_hours"] for g in groups] == [2.0, 2.0, 2.0]
    assert summarize(groups)["hours"] == 6.1
    assert by_project(groups)["TECH-OPS: TS20"]["hours"] == 6.1
